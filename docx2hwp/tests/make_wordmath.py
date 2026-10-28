# Word 가 실제로 쓰는 형태의 OMML(eqArr 정렬, lim, 적분, 행렬, 한글 텍스트 등)을 넣은 docx
import docx
from docx.oxml import parse_xml
M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
RPR = '<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr>'
def r(t, sty=None, nor=False):
    mr = ''
    if sty or nor:
        mr = '<m:rPr>' + ('<m:nor/>' if nor else '') + (f'<m:sty m:val="{sty}"/>' if sty else '') + '</m:rPr>'
    return f'<m:r>{mr}{RPR}<m:t xml:space="preserve">{t}</m:t></m:r>'
eqs = {
 "연립(정렬)": '<m:oMathPara><m:oMath><m:eqArr><m:e>' + r('2x+y') + r('&amp;=5') + '</m:e><m:e>' + r('x-y') + r('&amp;=1') + '</m:e></m:eqArr></m:oMath></m:oMathPara>',
 "극한": '<m:oMath><m:func><m:fName><m:limLow><m:e>' + r('lim','p') + '</m:e><m:lim>' + r('n→∞') + '</m:lim></m:limLow></m:fName><m:e><m:sSup><m:e><m:d><m:e>' + r('1+') + '<m:f><m:num>' + r('1') + '</m:num><m:den>' + r('n') + '</m:den></m:f></m:e></m:d></m:e><m:sup>' + r('n') + '</m:sup></m:sSup></m:e></m:func>' + r('=e') + '</m:oMath>',
 "적분": '<m:oMath><m:nary><m:naryPr><m:limLoc m:val="subSup"/></m:naryPr><m:sub>' + r('a') + '</m:sub><m:sup>' + r('b') + '</m:sup><m:e>' + r('f') + '<m:d><m:e>' + r('x') + '</m:e></m:d>' + r('dx') + '</m:e></m:nary>' + r('=F') + '<m:d><m:e>' + r('b') + '</m:e></m:d>' + r('-F') + '<m:d><m:e>' + r('a') + '</m:e></m:d></m:oMath>',
 "행렬식": '<m:oMath><m:d><m:dPr><m:begChr m:val="|"/><m:endChr m:val="|"/></m:dPr><m:e><m:m><m:mr><m:e>' + r('a') + '</m:e><m:e>' + r('b') + '</m:e></m:mr><m:mr><m:e>' + r('c') + '</m:e><m:e>' + r('d') + '</m:e></m:mr></m:m></m:e></m:d>' + r('=ad-bc') + '</m:oMath>',
 "한글텍스트": '<m:oMath>' + r('넓이','p',True) + r('=') + '<m:f><m:num>' + r('1') + '</m:num><m:den>' + r('2') + '</m:den></m:f>' + r('×') + r('밑변','p',True) + r('×') + r('높이','p',True) + '</m:oMath>',
 "삼각함수": '<m:oMath><m:func><m:fName><m:sSup><m:e>' + r('sin','p') + '</m:e><m:sup>' + r('2') + '</m:sup></m:sSup></m:fName><m:e>' + r('θ') + '</m:e></m:func>' + r('+') + '<m:func><m:fName><m:sSup><m:e>' + r('cos','p') + '</m:e><m:sup>' + r('2') + '</m:sup></m:sSup></m:fName><m:e>' + r('θ') + '</m:e></m:func>' + r('=1') + '</m:oMath>',
 "조합/벡터": '<m:oMath><m:sSub><m:e>' + r('C') + '</m:e><m:sub>' + r('n,r') + '</m:sub></m:sSub>' + r('=') + '<m:f><m:num>' + r('n!') + '</m:num><m:den>' + r('r!') + '<m:d><m:e>' + r('n-r') + '</m:e></m:d>' + r('!') + '</m:den></m:f>' + r(',  ') + '<m:acc><m:accPr><m:chr m:val="⃗"/></m:accPr><m:e>' + r('AB') + '</m:e></m:acc>' + r('·') + '<m:acc><m:accPr><m:chr m:val="⃗"/></m:accPr><m:e>' + r('CD') + '</m:e></m:acc>' + r('=0') + '</m:oMath>',
 "위아래묶음": '<m:oMath><m:limLow><m:e><m:groupChr><m:e>' + r('1+2+⋯+n') + '</m:e></m:groupChr></m:e><m:lim>' + r('n개') + '</m:lim></m:limLow></m:oMath>',
 "막대·켤레": '<m:oMath><m:bar><m:barPr><m:pos m:val="top"/></m:barPr><m:e>' + r('z') + '</m:e></m:bar>' + r('=a-bi, ') + '<m:sSubSup><m:e>' + r('x') + '</m:e><m:sub>' + r('n') + '</m:sub><m:sup>' + r('′') + '</m:sup></m:sSubSup>' + r('≠') + '<m:rad><m:deg>' + r('n') + '</m:deg><m:e>' + r('x') + '</m:e></m:rad></m:oMath>',
}
d = docx.Document()
d.add_heading('Word 수식 테스트', 1)
for name, xml in eqs.items():
    p = d.add_paragraph(name + ': ')
    p._p.append(parse_xml(xml.replace('<m:oMath', f'<m:oMath {M}', 1) if xml.startswith('<m:oMath>') else xml.replace('<m:oMathPara>', f'<m:oMathPara {M}>', 1)))
d.save('tests/wordmath.docx')
